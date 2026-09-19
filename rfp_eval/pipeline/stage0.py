"""Stage 0 -- turn the original RFP into a canonical question list.

Heuristics run first and cost nothing. They handle the two common shapes: an
Excel questionnaire with a question column and a section column, and a PDF with
numbered items. When they come up short, one LLM call normalizes the document
instead.

Either way the result lands in the UI for review before a single vendor call is
spent, because every downstream stage keys off these question ids.
"""

from __future__ import annotations

import re
from pathlib import Path

from ..llm.base import LLMProvider
from ..models import Question, Questionnaire
from ..parsing import Document
from ..prompts import load_prompt
from ..schema_utils import strict_schema

# Below this, the heuristics have not really found a questionnaire.
MIN_HEURISTIC_QUESTIONS = 3

# A question is prose, not a label.
MIN_QUESTION_CHARS = 20

_QUESTION_HEADERS = re.compile(
    r"question|requirement|query|item|ask|description|criteria|criterion", re.I
)
_CATEGORY_HEADERS = re.compile(r"categor|section|area|domain|theme|topic|group", re.I)
_REQUIRED_HEADERS = re.compile(r"required|mandatory|must", re.I)

# "1.", "1.2", "Q3", "Q3:", "3)" at the start of a line.
_NUMBERED = re.compile(
    r"^\s*(?:(?P<qid>Q\s?\d+(?:\.\d+)*)|(?P<num>\d+(?:\.\d+)*))\s*[.):]?\s+(?P<text>\S.*)$"
)
# A short bold-ish line with no terminal punctuation reads as a section heading.
_HEADING = re.compile(r"^\s*(?:[A-Z][A-Za-z0-9 &/,'-]{2,60})\s*:?\s*$")


# ---------------------------------------------------------------------------
# Heuristics
# ---------------------------------------------------------------------------


def _heuristic_excel(path: Path) -> list[Question]:
    """Find the question column and, if present, the category column."""
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception:
        return []

    questions: list[Question] = []
    counter = 0

    try:
        for sheet in workbook.worksheets:
            rows = [r for r in sheet.iter_rows(values_only=True) if r and any(r)]
            if len(rows) < 2:
                continue

            headers = [str(c).strip() if c is not None else "" for c in rows[0]]
            body = rows[1:]

            q_col = _pick_question_column(headers, body)
            if q_col is None:
                continue

            c_col = _pick_column(headers, _CATEGORY_HEADERS)
            r_col = _pick_column(headers, _REQUIRED_HEADERS)
            # With no category column, the sheet name is the next best grouping.
            default_category = sheet.title

            for row in body:
                text = _cell(row, q_col)
                if len(text) < MIN_QUESTION_CHARS:
                    continue
                counter += 1
                category = _cell(row, c_col) or default_category
                required_raw = _cell(row, r_col).lower()
                required = required_raw not in ("no", "n", "false", "optional", "0")
                questions.append(
                    Question(
                        id=f"Q{counter}",
                        category=category.strip() or "General",
                        text=text,
                        required=required,
                    )
                )
    finally:
        workbook.close()

    return questions


def _cell(row: tuple, index: int | None) -> str:
    if index is None or index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _pick_column(headers: list[str], pattern: re.Pattern) -> int | None:
    for i, header in enumerate(headers):
        if header and pattern.search(header):
            return i
    return None


def _pick_question_column(headers: list[str], body: list[tuple]) -> int | None:
    """Prefer a question-ish header; otherwise the column with the longest prose."""
    by_header = _pick_column(headers, _QUESTION_HEADERS)
    if by_header is not None:
        return by_header

    width = max((len(r) for r in body), default=0)
    best_col, best_len = None, 0.0
    for col in range(width):
        values = [_cell(r, col) for r in body]
        filled = [v for v in values if v]
        if len(filled) < MIN_HEURISTIC_QUESTIONS:
            continue
        mean_len = sum(len(v) for v in filled) / len(filled)
        if mean_len > best_len:
            best_col, best_len = col, mean_len

    # A column of short labels is a category, not a question.
    return best_col if best_len >= 30 else None


def _heuristic_pdf(document: Document) -> list[Question]:
    """Pull numbered items out, tracking the most recent heading as category."""
    questions: list[Question] = []
    category = "General"
    counter = 0

    for unit in document.units:
        for line in unit.text.splitlines():
            stripped = line.strip()
            if not stripped:
                continue

            match = _NUMBERED.match(stripped)
            if match:
                text = match.group("text").strip()
                if len(text) < MIN_QUESTION_CHARS:
                    continue
                counter += 1
                qid = (match.group("qid") or "").replace(" ", "") or f"Q{match.group('num')}"
                questions.append(
                    Question(id=qid or f"Q{counter}", category=category, text=text)
                )
                continue

            # Track section headings so questions inherit a meaningful category.
            if len(stripped) <= 60 and _HEADING.match(stripped) and not stripped.endswith("."):
                category = stripped.rstrip(":").strip() or category

    return _dedupe_ids(questions)


def _dedupe_ids(questions: list[Question]) -> list[Question]:
    """Document numbering restarts per section; ids must stay unique."""
    seen: dict[str, int] = {}
    for q in questions:
        if q.id in seen:
            seen[q.id] += 1
            q.id = f"{q.id}-{seen[q.id]}"
        else:
            seen[q.id] = 0
    return questions


def heuristic_questionnaire(document: Document, path: Path | None) -> Questionnaire | None:
    """Try to parse without an LLM. Returns None when the result is unconvincing."""
    if document.kind == "excel" and path is not None:
        questions = _heuristic_excel(path)
    elif document.kind == "pdf":
        questions = _heuristic_pdf(document)
    else:
        questions = []

    if len(questions) < MIN_HEURISTIC_QUESTIONS:
        return None

    return Questionnaire(
        rfp_name=Path(document.filename).stem,
        questions=questions,
        source_file=document.filename,
        parsed_by="heuristic",
    )


# ---------------------------------------------------------------------------
# LLM normalization
# ---------------------------------------------------------------------------


def llm_questionnaire(
    document: Document,
    provider: LLMProvider,
    *,
    model: str,
    max_output_tokens: int,
    max_chunk_tokens: int,
) -> Questionnaire:
    """Normalize the RFP into canonical questions.

    A long RFP is chunked and each chunk normalized, then the question lists are
    concatenated -- same no-truncation guarantee Stage 1 gives vendor documents.
    Most questionnaires are one chunk and cost exactly one call.
    """
    from .chunk_helpers import chunks_for

    prompt = load_prompt("stage0_questionnaire")
    schema = strict_schema(Questionnaire, exclude={"source_file", "parsed_by"})
    chunks = chunks_for(document, max_chunk_tokens, provider=provider, model=model)

    collected: list[Question] = []
    rfp_name = ""

    for chunk in chunks:
        system, user = prompt.render(
            document_text=chunk.text,
            filename=document.filename,
        )
        payload = provider.complete_json(
            system=system,
            user=user,
            schema=schema,
            model=model,
            max_tokens=max_output_tokens,
            stage="stage0",
            cache_system=True,
        )
        parsed = Questionnaire.model_validate(payload)
        rfp_name = rfp_name or parsed.rfp_name
        collected.extend(parsed.questions)

    return Questionnaire(
        rfp_name=rfp_name or Path(document.filename).stem,
        questions=_dedupe_ids(_dedupe_text(collected)),
        source_file=document.filename,
        parsed_by="llm",
    )


def _dedupe_text(questions: list[Question]) -> list[Question]:
    """Drop questions repeated across chunk boundaries."""
    seen: set[str] = set()
    unique: list[Question] = []
    for q in questions:
        key = q.text.strip().lower()
        if key in seen:
            continue
        seen.add(key)
        unique.append(q)
    return unique


def build_questionnaire(
    document: Document,
    provider: LLMProvider,
    *,
    path: Path | None,
    model: str,
    max_output_tokens: int,
    max_chunk_tokens: int,
    force_llm: bool = False,
) -> Questionnaire:
    """Heuristics first, LLM as the fallback."""
    if not force_llm:
        parsed = heuristic_questionnaire(document, path)
        if parsed is not None:
            return parsed

    return llm_questionnaire(
        document,
        provider,
        model=model,
        max_output_tokens=max_output_tokens,
        max_chunk_tokens=max_chunk_tokens,
    )
