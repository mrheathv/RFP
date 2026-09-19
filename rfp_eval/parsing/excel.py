"""Excel parsing via openpyxl.

Each non-empty row becomes one ``TextUnit``, serialized as ``Header: value``
pairs so the model sees which column a value came from. A row is the atomic
unit -- the chunker will never split one, which keeps a question and its answer
together even in a workbook that runs to thousands of rows.
"""

from __future__ import annotations

from pathlib import Path

from ..errors import ParseError
from . import Document, TextUnit

# A header row with more blanks than labels is not really a header.
_HEADER_MIN_FILL = 0.5
# Headers are labels ("Question", "Section"), not prose. A sheet with no header
# row must not have its first data row eaten as one, which would drop a real
# question and mislabel every column below it.
_HEADER_MAX_CHARS = 50
_HEADER_MAX_WORDS = 6


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _looks_like_header(row: tuple) -> bool:
    values = [_cell_str(v) for v in row]
    filled = [v for v in values if v]
    if not filled:
        return False
    if len(filled) / len(values) < _HEADER_MIN_FILL:
        return False
    return all(
        len(v) <= _HEADER_MAX_CHARS
        and len(v.split()) <= _HEADER_MAX_WORDS
        and not v.rstrip().endswith((".", "?", "!"))
        for v in filled
    )


def parse_excel(path: Path, display_name: str) -> Document:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ParseError(display_name, f"openpyxl is not installed: {exc}") from exc

    units: list[TextUnit] = []
    sheet_names: list[str] = []

    try:
        # data_only=True gives computed values rather than formula strings.
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError(
            display_name,
            f"could not read workbook ({type(exc).__name__}: {exc}). "
            "It may be corrupt, password-protected, or not a real .xlsx file.",
        ) from exc

    try:
        for sheet in workbook.worksheets:
            sheet_names.append(sheet.title)
            headers: list[str] = []

            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row is None:
                    continue
                values = [_cell_str(v) for v in row]
                if not any(values):
                    continue

                if not headers and _looks_like_header(row):
                    headers = values
                    continue

                # Serialize as Header: value so column meaning survives.
                parts = []
                for col_index, value in enumerate(values):
                    if not value:
                        continue
                    header = headers[col_index] if col_index < len(headers) else ""
                    parts.append(f"{header}: {value}" if header else value)

                if not parts:
                    continue

                units.append(
                    TextUnit(
                        text=" | ".join(parts),
                        locator=f"sheet '{sheet.title}', row {row_index}",
                        section=sheet.title,
                    )
                )
    finally:
        workbook.close()

    return Document(
        filename=display_name,
        units=units,
        kind="excel",
        meta={"sheets": sheet_names},
    )
